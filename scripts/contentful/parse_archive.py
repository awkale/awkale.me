#!/usr/bin/env python3
"""Parse 'Wikipedia BSO Archive.xlsx' into a normalized entity graph for Contentful.

Emits bso-graph.json: entity tables keyed by deterministic Contentful entry id,
plus a report of anything the parser could not confidently interpret.

Sheet conventions this relies on (verified against the source):
  - 'SEASON n' in col A starts a season; 'SEASON n, cont.' continues it.
  - A non-empty col A that isn't a season header starts a new concert.
  - Leading whitespace in the Piece cell means the text continues the item above.
  - A row with no Piece and no Composer carries extra soloist credits for the
    item above.
"""
import json, re, sys, hashlib, unicodedata, datetime, collections
from pathlib import Path

import openpyxl

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "Wikipedia BSO Archive.xlsx")
OUT = Path(__file__).parent / "bso-graph.json"

# ---------------------------------------------------------------- helpers

def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s)
                   if not unicodedata.combining(c))

def slugify(s, maxlen=40):
    s = strip_accents(str(s)).lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:maxlen].strip("-") or "x"

def h6(s):
    return hashlib.sha1(str(s).encode("utf-8")).hexdigest()[:6]

def norm(s):
    """Aggressive normalization for duplicate detection."""
    s = strip_accents(str(s)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()

PARTICLES = {"van", "von", "de", "di", "du", "del", "della", "le", "la", "der", "den", "ten"}

def name_key(first, last):
    """Match key that ignores nobiliary particles, so 'van Beethoven, Ludwig'
    and an existing entry stored as first='Ludwig' last='Beethoven' collide."""
    toks = [t for t in norm(f"{first} {last}").split() if t and t not in PARTICLES]
    return " ".join(sorted(toks))

BLANK = {"", "unknown", "none", "n/a", "na", "various", "tbd"}
def is_blank(v):
    return v is None or str(v).strip().lower() in BLANK

# instrument/voice/role values allowed by the soloist.instrument enum
ENUM = ["Violin","Viola","Violoncello","Double Bass","Flute","Oboe","Clarinet","Bassoon",
        "French Horn","Trumpet","Trombone","Tuba","Percussion","Harp","Piano","Saxophone",
        "Alto Saxophone","Harpsichord","Guitar","Accordion","Marimba","Vibraphone","Xylophone",
        "Timpani","Drums","Basso Continuo","Soprano","Mezzo-Soprano","Contralto","Alto",
        "Tenor","Baritone","Bass","Director","Narrator","Soloist"]
# sheet spelling -> canonical enum value (keeps us off duplicate instruments)
ALIAS = {"cello": "Violoncello", "violincello": "Violoncello", "horn": "French Horn",
         "accordian": "Accordion", "french horn": "French Horn", "contrabass": "Double Bass",
         "mezzo soprano": "Mezzo-Soprano", "mezzo-soprano": "Mezzo-Soprano"}
ENUM_LOOKUP = {norm(v): v for v in ENUM}
ENUM_LOOKUP.update({norm(k): v for k, v in ALIAS.items()})

ENSEMBLE_RX = re.compile(
    r"chorus|choir|chorale|choralettes|ensemble|quintet|quartet|sextet|trio|octet|"
    r"orchestra|philharmonia|band|players|singers|company|community sing|society",
    re.I)
# groups whose names carry none of the keywords above
ENSEMBLE_NAMES = {"spiritus et anima", "the nancy beth falloon waltzers"}

def ensemble_kind(name):
    n = name.lower()
    if "chorale" in n or "chorus" in n or "choral" in n: return "Chorus"
    if "choir" in n: return "Choir"
    if "brass" in n: return "Brass Ensemble"
    if "percussion" in n: return "Percussion Ensemble"
    if re.search(r"quintet|quartet|sextet|trio|octet|ensemble|players", n): return "Chamber Ensemble"
    if "singers" in n: return "Vocal Group"
    return "Other"

# ordered: more specific patterns first
GENRE_RX = [
    ("Concerto Grosso", r"\bconcerto grosso\b"),
    ("Symphony",        r"\bsymphon(y|ie|ia)\b|\bsinfoni"),
    ("Concerto",        r"\bconcerto\b|\bconcertino\b"),
    ("Overture",        r"\bovertur"),
    ("Suite",           r"\bsuite\b"),
    ("Mass",            r"\bmass\b|\brequiem\b|\bte deum\b|\bmagnificat\b"),
    ("Cantata",         r"\bcantata\b|\boratorio\b"),
    ("Variations",      r"\bvariations\b"),
    ("Rhapsody",        r"\brhapsod"),
    ("Fantasia",        r"\bfantas"),
    ("Serenade",        r"\bserenade\b|\bdivertimento\b"),
    ("Ballet",          r"\bballet\b"),
    ("Tone Poem",       r"\btone poem\b|\bsymphonic poem\b"),
    ("Prelude",         r"\bprelude\b|\bvorspiel\b"),
    ("Waltz",           r"\bwaltz\b|\bvalse\b"),
    ("March",           r"\bmarch\b"),
    ("Sonata",          r"\bsonata\b"),
    ("Aria",            r"\baria\b|,\s*from\b|\bfrom (the |la |der |das |le )?\w"),
]

KEY_RX = re.compile(r"\bin\s+([A-G](?:-flat|-sharp|b|#)?\s+(?:Major|Minor))\b", re.I)
NICK_RX = re.compile(r'\("([^"]+)"\)')
NOTE_RX = re.compile(r"world premiere|premiere|concert version|excerpts?|\bmvts?\b|"
                     r"\bmovement\b|abridged|arr\.|orch\.", re.I)
MONTHS = {m.lower(): i for i, m in enumerate(
    ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"], 1)}

def parse_date(v):
    """-> (iso_date_or_None, note_or_None)"""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat(), None
    if isinstance(v, datetime.date):
        return v.isoformat(), None
    s = str(v).strip()
    m = re.search(r"([A-Za-z]{3})\w*\.?\s+(\d{1,2}),?\s+(\d{4})", s)
    if m and m.group(1).lower() in MONTHS:
        try:
            return datetime.date(int(m.group(3)), MONTHS[m.group(1).lower()],
                                 int(m.group(2))).isoformat(), None
        except ValueError:
            pass
    return None, s   # unparseable -> keep verbatim in dateNote

# ---------------------------------------------------------------- id registry

class Registry:
    """Allocates stable, collision-free entry ids and dedupes by match key."""
    def __init__(self):
        self.tables = collections.defaultdict(dict)   # type -> {id: fields}
        self.by_key = {}                              # (type, key) -> id
        self.used = set()

    def _uniq(self, base):
        cid, n = base, 2
        while cid in self.used:
            cid = f"{base[:60]}-{n}"; n += 1
        self.used.add(cid)
        return cid

    def get(self, ctype, key, id_base, build):
        """Return existing id for key, else allocate one and build the fields."""
        k = (ctype, key)
        if k in self.by_key:
            return self.by_key[k]
        cid = self._uniq(id_base)
        self.by_key[k] = cid
        self.tables[ctype][cid] = build()
        return cid

R = Registry()
report = collections.defaultdict(list)

# ---------------------------------------------------------------- entity builders

def get_composer(raw):
    """'van Beethoven, Ludwig' -> composer id. Returns None for blanks."""
    s = re.sub(r"\s+", " ", str(raw)).strip().rstrip(",")
    if is_blank(s):
        return None
    if "," in s:
        last, first = s.split(",", 1)
        last, first = last.strip(), first.strip()
    else:
        last, first = s, ""          # 'Traditional', 'English Carol (arr. ...)'
    return R.get("composer", name_key(first, last),
                 f"cmp-{slugify(last + '-' + first)}",
                 lambda: {"firstName": first or None, "lastName": last,
                          "sortName": s, "dateOfBirth": None, "dateOfDeath": None})

def get_conductor(raw):
    s = re.sub(r"\s+", " ", str(raw)).strip()
    if is_blank(s):
        return None
    parts = s.split()
    first, last = " ".join(parts[:-1]), parts[-1]
    return R.get("conductor", name_key(first, last), f"cnd-{slugify(s)}",
                 lambda: {"firstName": first or None, "lastName": last})

ORCH_NAMES = {"BSO": "Brooklyn Symphony Orchestra",
              "BHO": "Brooklyn Heights Orchestra",
              "BHMS": "Brooklyn Heights Music Society"}

def get_orchestra(raw):
    if is_blank(raw):
        return None
    ab = str(raw).strip().upper()
    name = ORCH_NAMES.get(ab, str(raw).strip())
    return R.get("orchestra", norm(name), f"orc-{slugify(ab)}",
                 lambda: {"name": name, "abbreviation": ab})

VENUE_FIX = {"brooklyn musuem of art": "Brooklyn Museum of Art"}

def get_hall(raw):
    if is_blank(raw) or "not available" in str(raw).lower():
        return None
    s = re.sub(r"\s+", " ", str(raw)).strip()
    s = VENUE_FIX.get(norm(s), s)
    # existing entries follow 'Walt Whitman Hall' + location 'Brooklyn College'
    if "," in s:
        name, loc = s.split(",", 1)
        name, loc = name.strip(), loc.strip()
    else:
        name, loc = s, None
    return R.get("hall", norm(name), f"hal-{slugify(name)}",
                 lambda: {"name": name, "location": loc, "slug": slugify(name, 60)})

def get_genre(name):
    return R.get("genre", norm(name), f"gen-{slugify(name)}", lambda: {"name": name})

def get_season(num, notes):
    # `label` is filled in a later pass, once the concerts exist. It cannot be
    # built here: the label carries the season's YEARS, and the year is not a
    # function of the number -- season 48 is 2021-2022, not 2020-2021, because
    # the cancelled COVID season consumed no number. See label_seasons().
    return R.get("season", num, f"sea-{num}",
                 lambda: {"number": num, "label": None,
                          "notes": notes or None})

def get_work(title, composer_id, composer_raw):
    title = re.sub(r"\s+", " ", str(title)).strip()
    ckey = composer_id or "anon"
    key = (ckey, norm(title))
    def build():
        km = KEY_RX.search(title)
        nm = NICK_RX.search(title)
        genre_id = None
        for gname, rx in GENRE_RX:
            if re.search(rx, title, re.I):
                genre_id = get_genre(gname)
                break
        return {"title": title,
                "slug": f"{slugify(composer_raw or 'anon', 24)}--{slugify(title, 34)}-{h6(key)}",
                "musicalKey": km.group(1).title() if km else None,
                "nickname": nm.group(1) if nm else None,
                "composer": composer_id,
                "genre": genre_id,
                "movement": None}
    return R.get("work", key, f"wrk-{slugify(title, 30)}-{h6(key)}", build)

def get_performer(credit):
    """One credit string -> (performer_id, credit_kind, character_or_None).

    A credit is 'Name' or 'Name, Role[, Role...]'. Multiple roles mean one
    player covering several instruments ('Bill Utley, Tabla, Temple Blocks,
    Drums'). Roles that aren't instruments/voices are opera characters.
    """
    s = re.sub(r"\s+", " ", str(credit)).strip().rstrip(",;")
    if is_blank(s) or s.endswith(":"):
        return None, None, None
    if s.lower().startswith("with "):          # stage direction, not a performer
        return None, None, None

    segs = [p.strip() for p in s.split(",") if p.strip()]
    name, roles = segs[0], segs[1:]

    if ENSEMBLE_RX.search(name) or norm(name) in ENSEMBLE_NAMES:
        return (R.get("ensemble", norm(name), f"ens-{slugify(name)}",
                      lambda: {"name": name, "kind": ensemble_kind(name)}),
                "ensemble", None)
    if not re.match(r"^[A-Z]", name) or len(name.split()) > 5:
        return None, None, None                 # not a personal name

    enum_vals = [ENUM_LOOKUP[norm(r)] for r in roles if norm(r) in ENUM_LOOKUP]
    others = [r for r in roles if norm(r) not in ENUM_LOOKUP]
    parts = name.split()
    first, last = " ".join(parts[:-1]), parts[-1]
    pid = R.get("soloist", name_key(first, last), f"sol-{slugify(name)}",
                lambda: {"firstName": first or None, "lastName": last,
                         "fullName": name, "instrument": None})
    # a person seen again with new instruments accumulates them
    rec = R.tables["soloist"][pid]
    if enum_vals:
        rec["instrument"] = sorted(set((rec.get("instrument") or []) + enum_vals))
    return pid, "soloist", (others[0] if others else None)

def split_credits(raw):
    """The sheet puts exactly one credit per cell -- verified: no cell uses
    ';', ' and ', or a newline as a separator, and the only ' & ' is inside
    the ensemble name 'Grace & Spiritus Chorale of Brooklyn'."""
    s = str(raw).strip()
    return [s] if s else []

# ---------------------------------------------------------------- main sweep

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["ARCHIVE"]

concerts = []          # ordered list of concert dicts
season_num, season_notes = None, None
cur = None             # current concert
cur_item = None        # current program item

sheet = [list(r) for r in ws.iter_rows(min_row=6, values_only=True)]

# ------------------------------------------------- source corrections (AWK-38)
#
# Three transcription errors in the spreadsheet, fixed HERE and keyed by sheet row
# rather than fixed in the file. `Wikipedia BSO Archive.xlsx` is a received primary
# source and this repo keeps such things intact -- see docs/archive/ for the others.
# A correction in code is greppable, shows up in a diff, and is covered by
# archive-corrections.test.ts; an edited binary is none of those.
#
# Every entry pins the value it EXPECTS to find, and the run ABORTS on a mismatch.
# A correction applied silently to data that has since changed is worse than no
# correction at all -- the same posture migrate_schema.py takes toward drift.
COL = {"date": 0, "piece": 1, "soloist": 2, "composer": 3,
       "conductor": 4, "orchestra": 5, "venue": 6}

SOURCE_CORRECTIONS = {
    888: {
        "why": "2007-12-16 left conductor and orchestra blank. Every neighbouring "
               "concert in the season reads Armstrong/BSO and the venue matches. "
               "ADR-0006 ships conductor as one of only two browse filters, so a "
               "blank made this the one played concert no filter could reach.",
        "expect": {"conductor": None, "orchestra": None},
        "set": {"conductor": "Nicholas Armstrong", "orchestra": "BSO"},
    },
    912: {
        "why": "Sat 2008-12-13 at Grand Street -- the FIRST night of a two-venue "
               "run, and the archive's only one. Its piece cell restates just the "
               "opening work while row 913 carries the full program, so this row "
               "contributes the occasion and not the repertoire. Piece and composer "
               "are cleared to stop it emitting a phantom one-work concert.",
        # Piece and composer are pinned as well as the venue, because this entry
        # CLEARS them: if the sheet is ever amended so this row carries repertoire
        # unique to the Grand Street night, deleting it silently would leave the
        # Saturday sharing the Sunday's program and lose the difference.
        "expect": {"venue": "Grand Street Campus High Schools, Brooklyn",
                   "piece": "Capriccio Italien",
                   "composer": "Tchaikovsky, Pyotr Ilyich"},
        "set": {"piece": None, "composer": None},
        "role": "run-first-night",
    },
    913: {
        "why": "Labelled 'Sun, Dec 13, 2008', but Dec 13 2008 was a SATURDAY. The "
               "weekday is the half that is right and the day-of-month is the typo. "
               "This is the Sunday at St Ann, sharing the Saturday's program. It "
               "keeps its OWN hall: get_hall(ven) already takes precedence over the "
               "inherited one, which is why a two-venue run needs no change to the "
               "hall logic -- only to the run DETECTION below.",
        "expect": {"date": "2008-12-13"},
        "set": {"date": datetime.date(2008, 12, 14)},
        "role": "run-continuation",
    },
}

RUN_FIRST_NIGHT = {rn for rn, c in SOURCE_CORRECTIONS.items()
                   if c.get("role") == "run-first-night"}
RUN_CONTINUATION = {rn for rn, c in SOURCE_CORRECTIONS.items()
                    if c.get("role") == "run-continuation"}


def apply_source_corrections():
    """Mutate `sheet` in place, refusing to touch a row that no longer matches."""
    for rn, c in sorted(SOURCE_CORRECTIONS.items()):
        i = rn - 6
        if not 0 <= i < len(sheet):
            sys.exit(f"correction row {rn} is outside the sheet "
                     f"({len(sheet)} rows read from row 6)")
        row = sheet[i]
        for field, want in c.get("expect", {}).items():
            got = row[COL[field]]
            got = parse_date(got)[0] if field == "date" else got
            ok = is_blank(got) if want is None else (got == want)
            if not ok:
                sys.exit(f"correction row {rn}: expected {field}={want!r}, found "
                         f"{got!r}. The source has changed, so this correction no "
                         f"longer describes it -- re-verify against the sheet "
                         f"before editing SOURCE_CORRECTIONS.")
        for field, value in c.get("set", {}).items():
            row[COL[field]] = value
        report["source_correction"].append(f"row {rn}: {c['why'].split('.')[0]}.")


apply_source_corrections()

def duplicate_header(i):
    """True if row i starts a concert that the very next row restates.

    MATCHES NOTHING as of AWK-38, and is kept on purpose. Rows 912-913 were the
    only instance -- the same date and the same opening piece, once as a date cell
    and once as text -- and they turned out not to be a duplicate at all, but a
    two-venue run whose Sunday had its day-of-month mistyped. SOURCE_CORRECTIONS
    fixes that date, the two dates now differ, and this no longer fires.

    Retained because the artifact it guards against is a property of how the sheet
    was maintained rather than of that one pair: honoring a genuine restatement
    emits a phantom concert holding only the first piece, silently. If an edit to
    the source reintroduces one, this catches and reports it.
    """
    a, piece = (list(sheet[i]) + [None] * 7)[:2]
    if a in (None, "") or i + 1 >= len(sheet):
        return False
    b, piece2 = (list(sheet[i + 1]) + [None] * 7)[:2]
    if b in (None, ""):
        return False
    if isinstance(b, str) and b.strip().upper().startswith("SEASON"):
        return False
    same_date = parse_date(a)[0] is not None and parse_date(a)[0] == parse_date(b)[0]
    same_piece = (str(piece or "").strip().lower() == str(piece2 or "").strip().lower()
                  and str(piece or "").strip() != "")
    return same_date and same_piece

for i, row in enumerate(sheet):
    rn = i + 6
    a, piece, sol, comp, cond, orch, ven = (list(row) + [None] * 7)[:7]
    if all(v in (None, "") for v in (a, piece, sol, comp, cond, orch, ven)):
        continue
    if duplicate_header(i):
        report["duplicate_concert_row"].append(f"row {rn}: {a!r} restated by row {rn+1}")
        continue

    # -- season header
    if isinstance(a, str) and a.strip().upper().startswith("SEASON"):
        m = re.match(r"SEASON\s+(\d+)", a.strip(), re.I)
        if m:
            n = int(m.group(1))
            extra = a.strip()[m.end():].strip(" ,-")
            if n != season_num:
                season_num, season_notes = n, (extra if extra.lower() != "cont." else None)
            elif extra and extra.lower() != "cont.":
                season_notes = ((season_notes + " ") if season_notes else "") + extra
            sid = get_season(season_num, season_notes)
            # a later header for the same season can carry the note (e.g. the
            # COVID cancellation on the second 'SEASON 47' row)
            if season_notes:
                R.tables["season"][sid]["notes"] = season_notes
        continue

    # -- new concert
    if a not in (None, ""):
        iso, note = parse_date(a)
        prev = cur
        # A dated row can be an additional performance of the preceding program
        # (a two-night run) rather than a new program. It then shares that
        # concert's program items, so the cast listed beneath it lands on the
        # right pieces -- and because "items" is shared by reference, any piece
        # on the row itself is appended to both concerts.
        #
        # The sheet writes such a row two ways:
        #   1. bare date, no piece and no composer (4 rows)
        #   2. date placed on the row of the run's *next* piece, leaving the
        #      conductor, orchestra and venue cells empty (4 rows)
        # Form 2 is only distinguishable by all three of conductor, orchestra
        # and venue being blank. Requiring all three matters: row 266
        # ('var. dates, 1983') has a blank conductor and venue but names BHO,
        # and is a genuine concert. Across all 1,484 rows these two forms match
        # exactly 8 rows, every one of them intended.
        continues_run = (is_blank(cond) and is_blank(orch) and is_blank(ven))
        bare_date = (is_blank(comp)
                     and not (piece is not None and str(piece).strip()))
        # A THIRD form, declared rather than sniffed: a run whose two nights are at
        # DIFFERENT venues. Neither test above can express it -- both require a
        # blank venue, and a blank venue is precisely what makes the second night
        # inherit the first's hall. The two conditions are the same condition, so
        # the pair is named in SOURCE_CORRECTIONS instead of being guessed at.
        #
        # The first night is excluded explicitly: its piece cell was cleared, which
        # would otherwise read as a `bare_date` and share the PREVIOUS concert's
        # program -- silently, and two months from the row it belongs to.
        declared_run = rn in RUN_CONTINUATION
        shares = (prev is not None
                  and rn not in RUN_FIRST_NIGHT
                  and (declared_run
                       or (bool(prev["items"]) and (bare_date or continues_run))))
        cur = {"date": iso, "dateNote": note, "season": season_num,
               "hall": get_hall(ven) or (prev["hall"] if shares else None),
               "orchestra": get_orchestra(orch) or (prev["orchestra"] if shares else None),
               "conductor": get_conductor(cond) or (prev["conductor"] if shares else None),
               "items": prev["items"] if shares else [],
               "shared": shares,
               "raw_date": str(a).strip(), "raw_venue": ven, "row": rn}
        if shares:
            cur["dateNote"] = note or f"Additional performance of the {prev['date'] or prev['raw_date']} program"
            # Empty on a declared run: the first night contributes the occasion and
            # this row carries the repertoire, so there is no item to continue onto.
            cur_item = prev["items"][-1] if prev["items"] else None
            how = ("declared two-venue run" if declared_run
                   else "bare date" if bare_date
                   else "date on the next piece's row")
            # `cur["raw_date"]` rather than `a`: a corrected date arrives here as a
            # datetime.date, whose repr would print as `datetime.date(2008, 12, 14)`
            # against every other line's `'Sun, Apr 26, 2015'`. This report is the
            # human-readable drift signal, so it stays uniform.
            report["shared_program"].append(
                f"row {rn}: {cur['raw_date']!r} shares the program of "
                f"{prev['date'] or prev['raw_date']!r} ({how})")
        else:
            cur_item = None
        concerts.append(cur)
        if note:
            report["unparseable_date"].append(f"row {rn}: {a!r}")
    if cur is None:
        report["orphan_row"].append(f"row {rn}: {piece!r}")
        continue

    piece_s = str(piece) if piece is not None else ""
    has_piece = piece_s.strip() != ""
    # leading whitespace + no composer  =>  wrapped continuation of the item above
    is_cont = has_piece and piece_s[:1].isspace() and is_blank(comp) and cur_item is not None

    if is_cont:
        cur_item["note"] = " ".join(filter(None, [cur_item.get("note"),
                                                  piece_s.strip()]))[:255]
    elif has_piece:
        t = piece_s.strip()
        if re.fullmatch(r"(works|records) not available", t, re.I):
            cur_item = {"label": t, "work": None, "composer": get_composer(comp) if not is_blank(comp) else None,
                        "soloists": [], "credits": [], "character": None, "note": t}
            cur["items"].append(cur_item)
        else:
            cid = get_composer(comp)
            wid = get_work(t, cid, comp)
            nm = NOTE_RX.search(t)
            cur_item = {"label": t[:255], "work": wid, "composer": None, "soloists": [],
                        "credits": [], "character": None,
                        "note": nm.group(0) if nm else None}
            cur["items"].append(cur_item)
    elif not is_blank(comp):
        # composer named but no piece (Oct 26 1982 'works not available' block)
        cid = get_composer(comp)
        cur_item = {"label": f"{str(comp).strip()} — work not recorded"[:255],
                    "work": None, "composer": cid, "soloists": [], "credits": [],
                    "character": None, "note": "work not recorded"}
        cur["items"].append(cur_item)

    # -- soloist credits attach to the current item
    if sol is not None and str(sol).strip():
        if cur_item is None:
            report["credit_without_item"].append(f"row {rn}: {sol!r}")
        else:
            for credit in split_credits(sol):
                credit = credit.strip()
                cur_item["credits"].append(credit)
                pid, kind, char = get_performer(credit)
                if pid:
                    if pid not in cur_item["soloists"]:
                        cur_item["soloists"].append(pid)
                    # only meaningful for a single-performer item; a full opera
                    # cast keeps its per-singer roles in `credits` instead
                    if char:
                        cur_item["character"] = (char if len(cur_item["credits"]) == 1
                                                 else None)
                elif not is_blank(credit):
                    report["credit_not_linked"].append(f"row {rn}: {credit!r}")

# ---------------------------------------------------------------- finalize concerts

DOW = {0:"Mon",1:"Tue",2:"Wed",3:"Thu",4:"Fri",5:"Sat",6:"Sun"}
seen_cid = collections.Counter()
for c in concerts:
    if c["date"]:
        base = "cnc-" + c["date"].replace("-", "")
    else:
        base = f"cnc-s{c['season']}-{slugify(c['raw_date'], 20)}"
    seen_cid[base] += 1
    cid = base if seen_cid[base] == 1 else f"{base}-{seen_cid[base]}"
    c["id"] = cid

    orch_ab = (R.tables["orchestra"][c["orchestra"]]["abbreviation"]
               if c["orchestra"] else None)
    label = c["date"] or c["raw_date"]
    c["title"] = " — ".join(filter(None, [label, orch_ab])) [:255]

    if c["shared"]:
        continue          # items already belong to (and are named by) the first date
    for i, it in enumerate(c["items"], start=1):
        it["order"] = i
        it["id"] = f"pi-{cid[4:]}-{i}"

# program items get their own table (shared items are written once)
for c in concerts:
    for it in c["items"]:
        R.tables["programItem"][it["id"]] = {
            k: it[k] for k in ("label","work","composer","soloists","credits",
                               "character","order","note")}

# ---------------------------------------------------------------- label seasons

def label_seasons():
    """`BSO Season <N>, <YYYY-YYYY>` for every season, years READ from concerts.

    AWK-59. Runs here rather than in get_season() because it needs the concert
    dates, which do not exist until parsing finishes.

    A season runs September to the following summer, so a concert in January
    belongs to the season that opened the previous autumn. Two seasons have no
    dated concert at all and take their year from season-orchestras.json, which
    is where hand-assigned values live with the reasoning that produced them."""
    decisions = json.loads((Path(__file__).parent / "season-orchestras.json").read_text())
    pattern = decisions["labelFormat"]["pattern"]
    hand = {k: v for k, v in decisions["handAssigned"].items() if k != "note"}

    first_date = {}
    for c in concerts:
        if not c["season"] or not c["date"]:
            continue
        sid = f"sea-{c['season']}"
        if sid not in first_date or c["date"] < first_date[sid]:
            first_date[sid] = c["date"]

    for sid, season in R.tables["season"].items():
        if sid in first_date:
            d = first_date[sid]
            start = int(d[:4]) if int(d[5:7]) >= 9 else int(d[:4]) - 1
            years = f"{start}-{start + 1}"
        elif sid in hand:
            years = hand[sid]["years"]
        else:
            # Fatal, not a report line. `label` is the season type's
            # displayField: leaving it None makes import_to_contentful.py skip
            # the field (build_fields drops None), which creates a season that
            # reads as untitled in every entry picker. The report is printed six
            # lines per key, so a soft warning here is one that gets missed.
            # backfill_seasons.derive() treats the same condition as fatal.
            sys.exit(
                f"{sid} has no dated concert and no entry in season-orchestras.json's "
                f"handAssigned. Add it there, with the reasoning that produced the year."
            )
        season["label"] = (pattern.replace("{institution}", "BSO")
                           .replace("{number}", str(season["number"]))
                           .replace("{years}", years))


label_seasons()

R.tables["concert"] = {c["id"]: {"title": c["title"], "date": c["date"],
                                 "dateNote": c["dateNote"],
                                 "season": f"sea-{c['season']}" if c["season"] else None,
                                 "hall": c["hall"], "orchestra": c["orchestra"],
                                 "conductor": c["conductor"],
                                 "program": [it["id"] for it in c["items"]]}
                       for c in concerts}

graph = {"types": {t: R.tables[t] for t in sorted(R.tables)},
         "report": {k: v for k, v in sorted(report.items())}}
OUT.write_text(json.dumps(graph, indent=1, ensure_ascii=False))

print(f"parsed {SRC.name}\n")
for t in sorted(R.tables):
    print(f"  {t:14s} {len(R.tables[t]):5d}")
print(f"\n  {'TOTAL':14s} {sum(len(v) for v in R.tables.values()):5d} entries")
print("\nreport:")
for k, v in sorted(report.items()):
    print(f"  {k}: {len(v)}")
    for line in v[:6]:
        print(f"     {line}")
print(f"\nwrote {OUT}")
